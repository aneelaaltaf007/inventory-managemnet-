# import libraries
import streamlit as st
from groq import Groq
from pypdf import PdfReader
from docx import Document
import io
import re
import uuid


# App Configuration

st.set_page_config(
    page_title="Document AI Assistant",
    page_icon="📚",
    layout="wide"
)


# Session State

if "groq_client" not in st.session_state:
    st.session_state.groq_client = None

if "documents" not in st.session_state:
    st.session_state.documents = {}

if "chunks" not in st.session_state:
    st.session_state.chunks = []

if "chats" not in st.session_state:
    st.session_state.chats = []

if "current_chat" not in st.session_state:
    st.session_state.current_chat = None

if "chat_messages" not in st.session_state:
    st.session_state.chat_messages = []


# Chat Functions

def save_current_chat():
    if not st.session_state.chat_messages:
        return

    title = "New Chat"

    for message in st.session_state.chat_messages:
        if message["role"] == "user":
            title = message["content"][:35]

            if len(message["content"]) > 35:
                title += "..."

            break

    chat_id = st.session_state.current_chat

    if chat_id:
        for chat in st.session_state.chats:
            if chat["id"] == chat_id:
                chat["messages"] = st.session_state.chat_messages.copy()
                chat["title"] = title
                return

    chat = {
        "id": str(uuid.uuid4()),
        "title": title,
        "messages": st.session_state.chat_messages.copy(),
        "pinned": False,
        "archived": False
    }

    st.session_state.chats.insert(0, chat)
    st.session_state.current_chat = chat["id"]


def load_chat(chat_id):
    for chat in st.session_state.chats:
        if chat["id"] == chat_id:
            st.session_state.current_chat = chat_id
            st.session_state.chat_messages = chat["messages"].copy()
            st.rerun()


# Document Extraction

def extract_pdf(file):
    text = ""

    try:
        reader = PdfReader(file)

        for page in reader.pages:
            page_text = page.extract_text()

            if page_text:
                text += page_text + "\n"

    except Exception as e:
        st.error(f"PDF reading error: {e}")

    return text


def extract_docx(file):
    text = ""

    try:
        document = Document(file)

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text += paragraph.text + "\n"

        for table in document.tables:
            for row in table.rows:
                row_text = [cell.text for cell in row.cells]
                text += " | ".join(row_text) + "\n"

    except Exception as e:
        st.error(f"DOCX reading error: {e}")

    return text


def extract_txt(file):
    try:
        return file.read().decode("utf-8", errors="ignore")

    except Exception as e:
        st.error(f"TXT reading error: {e}")
        return ""


def extract_document(file, file_name):
    file_name = file_name.lower()

    if file_name.endswith(".pdf"):
        return extract_pdf(file)

    if file_name.endswith(".docx"):
        return extract_docx(file)

    if file_name.endswith(".txt"):
        return extract_txt(file)

    return ""


# Text Processing

def clean_text(text):
    return re.sub(r"\s+", " ", text).strip()


def create_chunks(text, file_name, chunk_size=350):
    text = clean_text(text)

    if not text:
        return []

    words = text.split()
    chunks = []

    for start in range(0, len(words), chunk_size):
        chunk_text = " ".join(
            words[start:start + chunk_size]
        )

        chunks.append({
            "file_name": file_name,
            "text": chunk_text
        })

    return chunks


def find_relevant_chunks(question, chunks, max_chunks=2):
    if not chunks:
        return []

    question_words = set(
        re.findall(
            r"\b[a-zA-Z0-9]+\b",
            question.lower()
        )
    )

    scored_chunks = []

    for chunk in chunks:
        chunk_words = set(
            re.findall(
                r"\b[a-zA-Z0-9]+\b",
                chunk["text"].lower()
            )
        )

        score = len(
            question_words.intersection(chunk_words)
        )

        scored_chunks.append((score, chunk))

    scored_chunks.sort(
        key=lambda item: item[0],
        reverse=True
    )

    return [
        item[1]
        for item in scored_chunks[:max_chunks]
    ]


# Sidebar

with st.sidebar:
    st.header("🔑 Groq API Key")

    groq_api_key = st.text_input(
        "Groq API Key",
        type="password",
        placeholder="gsk_..."
    )

    if groq_api_key:
        try:
            st.session_state.groq_client = Groq(
                api_key=groq_api_key
            )

            st.success("API key connected ✅")

        except Exception as e:
            st.error(f"Invalid API key: {e}")

    st.divider()

    if st.button(
        "➕ New Chat",
        use_container_width=True
    ):
        save_current_chat()
        st.session_state.current_chat = None
        st.session_state.chat_messages = []
        st.rerun()

    st.subheader("💬 Chat History")

    pinned_chats = [
        chat
        for chat in st.session_state.chats
        if chat["pinned"] and not chat["archived"]
    ]

    normal_chats = [
        chat
        for chat in st.session_state.chats
        if not chat["pinned"] and not chat["archived"]
    ]

    archived_chats = [
        chat
        for chat in st.session_state.chats
        if chat["archived"]
    ]

    if pinned_chats:
        st.caption("📌 Pinned")

        for chat in pinned_chats:
            if st.button(
                chat["title"],
                key=f"open_pinned_{chat['id']}",
                use_container_width=True
            ):
                load_chat(chat["id"])

            menu = st.popover(
                "⋮",
                use_container_width=True
            )

            with menu:
                if st.button(
                    "Unpin",
                    key=f"unpin_{chat['id']}"
                ):
                    chat["pinned"] = False
                    st.rerun()

                if st.button(
                    "Archive",
                    key=f"archive_{chat['id']}"
                ):
                    chat["archived"] = True
                    st.rerun()

                if st.button(
                    "Delete",
                    key=f"delete_{chat['id']}"
                ):
                    st.session_state.chats.remove(chat)
                    st.rerun()

    if normal_chats:
        st.caption("Recent")

        for chat in normal_chats:
            if st.button(
                chat["title"],
                key=f"open_recent_{chat['id']}",
                use_container_width=True
            ):
                load_chat(chat["id"])

            menu = st.popover(
                "⋮",
                use_container_width=True
            )

            with menu:
                if st.button(
                    "Pin",
                    key=f"pin_{chat['id']}"
                ):
                    chat["pinned"] = True
                    st.rerun()

                if st.button(
                    "Archive",
                    key=f"archive_recent_{chat['id']}"
                ):
                    chat["archived"] = True
                    st.rerun()

                if st.button(
                    "Delete",
                    key=f"delete_recent_{chat['id']}"
                ):
                    st.session_state.chats.remove(chat)
                    st.rerun()

    if archived_chats:
        with st.expander("🗄️ Archived"):
            for chat in archived_chats:
                st.write(chat["title"])

                if st.button(
                    "Restore",
                    key=f"restore_{chat['id']}"
                ):
                    chat["archived"] = False
                    st.rerun()

                if st.button(
                    "Delete",
                    key=f"delete_archived_{chat['id']}"
                ):
                    st.session_state.chats.remove(chat)
                    st.rerun()


# Main Page

st.title("📚 Document AI Assistant")

st.write(
    "Upload your documents and ask questions about them using Groq AI."
)


# Select Documents

st.subheader("📄 Select Your Document")

uploaded_files = st.file_uploader(
    "Choose a PDF, Word document, or text file",
    type=["pdf", "docx", "txt"],
    accept_multiple_files=True,
    help="Select one or more PDF, DOCX, or TXT files from your computer."
)

if uploaded_files:
    st.write(
        f"📎 {len(uploaded_files)} file(s) selected"
    )

    for uploaded_file in uploaded_files:
        st.write(
            f"📄 {uploaded_file.name}"
        )

    if not groq_api_key:
        st.warning(
            "Please enter your Groq API key in the sidebar."
        )

    else:
        if st.button(
            "📤 Process Documents",
            use_container_width=True
        ):
            with st.spinner("Processing documents..."):
                new_documents = 0

                for uploaded_file in uploaded_files:
                    file_name = uploaded_file.name

                    if file_name in st.session_state.documents:
                        continue

                    try:
                        file_bytes = uploaded_file.getvalue()
                        file_copy = io.BytesIO(file_bytes)

                        text = extract_document(
                            file_copy,
                            file_name
                        )

                        if text.strip():
                            st.session_state.documents[file_name] = text

                            chunks = create_chunks(
                                text,
                                file_name,
                                chunk_size=350
                            )

                            st.session_state.chunks.extend(chunks)
                            new_documents += 1

                        else:
                            st.warning(
                                f"No readable text found in {file_name}."
                            )

                    except Exception as e:
                        st.error(
                            f"Error processing {file_name}: {e}"
                        )

                if new_documents:
                    st.success(
                        f"{new_documents} document(s) processed successfully! ✅"
                    )
                else:
                    st.info(
                        "These documents are already processed."
                    )


# Uploaded Documents

if st.session_state.documents:
    st.divider()

    st.subheader("📚 Your Documents")

    for index, file_name in enumerate(
        st.session_state.documents,
        start=1
    ):
        word_count = len(
            st.session_state.documents[file_name].split()
        )

        st.write(
            f"**{index}.** 📄 {file_name} — "
            f"{word_count:,} words"
        )

    st.success(
        f"📚 {len(st.session_state.documents)} "
        f"document(s) ready."
    )


# Chat Controls

col1, col2 = st.columns(2)

with col1:
    if st.button(
        "🆕 New Chat",
        use_container_width=True
    ):
        save_current_chat()
        st.session_state.current_chat = None
        st.session_state.chat_messages = []
        st.rerun()

with col2:
    if st.button(
        "🗑️ Clear Documents",
        use_container_width=True
    ):
        st.session_state.documents = {}
        st.session_state.chunks = []
        st.rerun()


# Chat Section

st.divider()

st.subheader("💬 Chat")

for message in st.session_state.chat_messages:
    with st.chat_message(message["role"]):
        st.markdown(message["content"])


question = st.chat_input(
    "Ask a question about your documents..."
)


# AI Response

if question:
    if not groq_api_key:
        st.warning(
            "Please enter your Groq API key in the sidebar."
        )
        st.stop()

    if not st.session_state.documents:
        st.warning(
            "Please select and process a document first."
        )
        st.stop()

    st.session_state.chat_messages.append({
        "role": "user",
        "content": question
    })

    with st.chat_message("user"):
        st.markdown(question)

    relevant_chunks = find_relevant_chunks(
        question,
        st.session_state.chunks,
        max_chunks=2
    )

    document_context = ""

    for chunk in relevant_chunks:
        document_context += (
            f"\n[{chunk['file_name']}]\n"
            f"{chunk['text']}\n"
        )

    system_message = """
You are a helpful Document AI Assistant.

Answer questions using the provided document context.

Do not invent information.

If the answer is not available in the documents, say:

"The information was not found in the uploaded documents."

Keep answers clear, useful and concise.
"""

    final_prompt = f"""
Document context:

{document_context}

Question:

{question}
"""

    messages = [
        {
            "role": "system",
            "content": system_message
        }
    ]

    messages.extend(
        st.session_state.chat_messages[-3:-1]
    )

    messages.append({
        "role": "user",
        "content": final_prompt
    })

    with st.chat_message("assistant"):
        with st.spinner("Thinking..."):
            try:
                client = st.session_state.groq_client

                response = client.chat.completions.create(
                    model="openai/gpt-oss-20b",
                    messages=messages,
                    temperature=0.2,
                    max_completion_tokens=600
                )

                answer = response.choices[0].message.content

                st.markdown(answer)

                st.session_state.chat_messages.append({
                    "role": "assistant",
                    "content": answer
                })

                save_current_chat()

            except Exception as e:
                st.error(
                    f"Groq API Error: {e}"
                )
